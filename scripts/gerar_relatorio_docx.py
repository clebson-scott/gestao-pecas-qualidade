#!/usr/bin/env python3
"""Gera RELATORIO_TECNICO.docx a partir do conteúdo da parte teórica,
formatado como um documento acadêmico (capa, títulos, tabelas, código).

Não é markdown->docx genérico: o conteúdo é escrito diretamente aqui para
garantir formatação limpa (fonte, espaçamento, tabelas) sem depender de
conversores externos (pandoc não está disponível no ambiente de execução).
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL_UNIFECAF = RGBColor(0x0B, 0x3C, 0x8A)
CINZA = RGBColor(0x44, 0x44, 0x44)

doc = Document()

# --- Estilo base ---
estilo_normal = doc.styles["Normal"]
estilo_normal.font.name = "Calibri"
estilo_normal.font.size = Pt(11)
estilo_normal.paragraph_format.space_after = Pt(8)
estilo_normal.paragraph_format.line_spacing = 1.15

for nome_estilo, tamanho, cor in [
    ("Heading 1", 18, AZUL_UNIFECAF),
    ("Heading 2", 14, AZUL_UNIFECAF),
    ("Heading 3", 12, CINZA),
]:
    estilo = doc.styles[nome_estilo]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(tamanho)
    estilo.font.color.rgb = cor
    estilo.font.bold = True


def sombrear_celula(celula, cor_hex="0B3C8A"):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), cor_hex)
    tcPr.append(shd)


def texto_branco_negrito(celula):
    for paragrafo in celula.paragraphs:
        for run in paragrafo.runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.bold = True


def paragrafo_codigo(texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(texto)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    # fundo levemente cinza no parágrafo inteiro
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    return p


def tabela_simples(cabecalhos, linhas, larguras=None):
    tabela = doc.add_table(rows=1, cols=len(cabecalhos))
    tabela.style = "Light Grid Accent 1"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tabela.rows[0].cells
    for i, titulo in enumerate(cabecalhos):
        hdr[i].text = titulo
        sombrear_celula(hdr[i])
        texto_branco_negrito(hdr[i])
    for linha in linhas:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha):
            celulas[i].text = str(valor)
    doc.add_paragraph()
    return tabela


# ============================== CAPA ==============================
doc.add_paragraph().add_run("\n" * 3)
titulo = doc.add_paragraph()
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = titulo.add_run("UniFECAF")
run.font.size = Pt(16)
run.font.bold = True
run.font.color.rgb = AZUL_UNIFECAF

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("Algoritmos e Lógica de Programação")
run.font.size = Pt(13)
run.font.color.rgb = CINZA

doc.add_paragraph().add_run("\n" * 4)

titulo2 = doc.add_paragraph()
titulo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = titulo2.add_run("Desafio de Automação Digital:\nGestão de Peças, Qualidade e Armazenamento")
run.font.size = Pt(22)
run.font.bold = True

doc.add_paragraph().add_run("\n")

subt = doc.add_paragraph()
subt.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subt.add_run("Parte Teórica — Análise e Discussão")
run.font.size = Pt(15)
run.italic = True
run.font.color.rgb = CINZA

doc.add_paragraph().add_run("\n" * 6)

autor = doc.add_paragraph()
autor.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = autor.add_run("Autor: Clebson Scott")
run.font.size = Pt(13)
run.font.bold = True

repo = doc.add_paragraph()
repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = repo.add_run("Repositório do código-fonte: github.com/clebson-scott/gestao-pecas-qualidade")
run.font.size = Pt(11)
run.font.color.rgb = CINZA

doc.add_page_break()

# ============================== SUMÁRIO (texto simples) ==============================
doc.add_heading("Sumário", level=1)
for item in [
    "1. Contextualização do desafio",
    "2. Como o raciocínio lógico foi estruturado",
    "3. Benefícios percebidos na solução",
    "4. Desafios enfrentados no desenvolvimento",
    "5. Reflexão final: expandindo o protótipo para um cenário real",
    "Referências",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_page_break()

# ============================== 1. CONTEXTUALIZAÇÃO ==============================
doc.add_heading("1. Contextualização do desafio", level=1)

doc.add_heading("1.1 Por que a automação é importante na indústria", level=2)
doc.add_paragraph(
    "O desafio proposto retrata uma situação extremamente comum em ambientes de "
    "manufatura de pequeno e médio porte: a inspeção de qualidade ainda feita "
    "manualmente, item a item, por um operador humano com régua, balança e uma "
    "tabela de referência de cores. Esse modelo carrega três fragilidades "
    "estruturais que a automação digital ataca diretamente:"
)
doc.add_paragraph(
    "1. Inconsistência de julgamento. Um operador humano, ao longo de um turno de "
    "oito horas, sofre fadiga visual e cognitiva. A régua e a balança não erram, "
    'mas a leitura e a decisão ("está dentro do critério?") sim. Um sistema '
    "automatizado aplica exatamente o mesmo critério, na milésima peça e na "
    "primeira, sem variação.",
    style="List Number",
)
doc.add_paragraph(
    "2. Custo de atraso. Cada peça inspecionada manualmente consome tempo de um "
    "profissional que poderia estar em uma atividade de maior valor agregado. Em "
    "escala, esse tempo se traduz diretamente em custo de operação e em gargalo "
    "de produção — a linha de montagem não pode andar mais rápido que a "
    "inspeção.",
    style="List Number",
)
doc.add_paragraph(
    "3. Ausência de rastreabilidade em tempo real. Quando o relatório de "
    "qualidade só é compilado ao final do turno (ou, pior, ao final da semana, "
    "numa planilha), qualquer problema sistemático na matéria-prima ou na "
    "calibração de uma máquina só é percebido depois que centenas de peças já "
    "foram produzidas fora do padrão.",
    style="List Number",
)
doc.add_paragraph(
    "A automação digital, mesmo em um protótipo simples como o deste desafio, "
    "ataca os três pontos simultaneamente: aplica um critério determinístico e "
    "auditável a cada peça, no instante do cadastro, e mantém um relatório "
    "sempre atualizado e disponível sob demanda — sem esperar o fim do turno."
)

doc.add_heading("1.2 O problema em termos de engenharia de software", level=2)
doc.add_paragraph(
    "Reduzido à sua essência lógica, o desafio pede três coisas encadeadas: uma "
    "função de classificação com múltiplos critérios (peso, cor, comprimento) "
    'que decide "aprovada" ou "reprovada" e explica o motivo; uma estrutura '
    "de agrupamento com capacidade limitada e fechamento automático (as caixas "
    "de 10 peças); e uma camada de agregação/relatório sobre os dados "
    "acumulados nos dois pontos anteriores. Essa decomposição — avaliação, "
    "armazenamento, relatório — é exatamente a que guiou a arquitetura do "
    "código, detalhada na seção 2."
)

# ============================== 2. RACIOCÍNIO LÓGICO ==============================
doc.add_heading("2. Como o raciocínio lógico foi estruturado", level=1)

doc.add_heading("2.1 Visão geral: por que dividir em módulos", level=2)
doc.add_paragraph(
    "A primeira decisão de projeto foi não escrever tudo em um único arquivo. "
    "Embora o desafio pudesse ser resolvido em um script único de 100-150 "
    "linhas, essa abordagem cresce mal: qualquer mudança na regra de peso "
    "obrigaria a reler todo o arquivo para achar onde ela está aplicada, e "
    "qualquer teste automatizado precisaria simular entrada de terminal para "
    "testar uma regra de negócio pura."
)
doc.add_paragraph(
    "A solução adotada separa o sistema em cinco responsabilidades isoladas, "
    "seguindo o princípio de responsabilidade única — cada módulo tem um único "
    "motivo para mudar:"
)
tabela_simples(
    ["Módulo", "Responsabilidade", "O que muda nele"],
    [
        ("models.py", "Estruturas de dados (Peça, Caixa, Enums)", "Novo atributo de peça exigido"),
        ("regras.py", "Critérios de aprovação/reprovação", "Fábrica muda a faixa de peso aceita"),
        (
            "estoque.py",
            "Orquestração (cadastro, remoção, relatório)",
            "Regra de fechamento de caixa muda",
        ),
        (
            "persistencia.py",
            "Salvar/carregar estado em disco",
            "Formato de armazenamento muda (ex.: banco de dados)",
        ),
        ("main.py", "Interface de terminal (menu)", "Forma de interação com o usuário muda"),
    ],
)

doc.add_heading("2.2 Decisões (estruturas condicionais)", level=2)
doc.add_paragraph(
    "A decisão central do sistema é a avaliação de qualidade "
    "(regras.py::avaliar_peca). Ao invés de uma cadeia de if/elif/else que "
    "interrompe na primeira condição verdadeira — abordagem comum, mas que "
    "esconde informação —, o sistema avalia todas as três condições "
    "independentemente:"
)
paragrafo_codigo(
    "verificacoes = (\n"
    "    avaliar_peso(peso),\n"
    "    avaliar_cor(cor),\n"
    "    avaliar_comprimento(comprimento),\n"
    ")\n"
    "motivos = tuple(m for m in verificacoes if m is not None)\n"
    "aprovada = len(motivos) == 0"
)
doc.add_paragraph(
    "Essa escolha lógica — testar todas as condições em vez de encadear com "
    "elif e parar na primeira falha — foi deliberada: uma peça que falha em "
    "peso e em cor precisa reportar os dois problemas, não só o primeiro. Do "
    "ponto de vista de quem opera a linha de produção, essa é a diferença entre "
    "um relatório útil (que aponta a causa raiz completa) e um relatório que "
    "obriga a testar a mesma peça repetidas vezes até descobrir todos os "
    "defeitos."
)
doc.add_paragraph(
    "Uma segunda decisão condicional relevante está no fechamento de caixas "
    "(models.py::Caixa.adicionar): a caixa só fecha depois de receber a peça "
    "que a completa, nunca antes — garantindo que a 10ª peça realmente entre na "
    "caixa corrente, e que a 11ª abra uma nova, sem nunca ultrapassar a "
    "capacidade."
)

doc.add_heading("2.3 Funções", level=2)
doc.add_paragraph(
    "Cada regra de qualidade é isolada em sua própria função pura — "
    "avaliar_peso, avaliar_cor, avaliar_comprimento — que recebe um valor e "
    "devolve None (aprovado) ou uma string explicando o motivo da reprovação. "
    '"Pura" aqui significa: a função não lê nem modifica nenhum estado '
    "externo, e sempre devolve o mesmo resultado para a mesma entrada. Essa "
    "propriedade é o que torna essas funções triviais de testar isoladamente e "
    "o que permite reaproveitá-las em qualquer contexto futuro — um endpoint de "
    "API, uma esteira física, outro protótipo — sem qualquer adaptação."
)

doc.add_heading("2.4 Condições e repetição no fluxo de armazenamento", level=2)
doc.add_paragraph(
    'O laço lógico "adicionar peça → caixa cheia? → fechar e abrir nova" é '
    "implementado sem nenhum loop explícito de repetição — cada chamada a "
    "cadastrar_peca() processa exatamente uma peça, e é o próprio menu em "
    "main.py que fornece a repetição, através de um laço while True: que "
    "mantém o sistema respondendo a comandos até o operador escolher "
    '"0. Sair".'
)
doc.add_paragraph(
    'Essa separação entre "a regra de armazenamento" (sem loop, uma peça por '
    'chamada) e "a repetição da interação com o usuário" (loop do menu) é '
    "outra aplicação do mesmo princípio de responsabilidade única: a lógica de "
    "armazenamento não sabe nem precisa saber que está sendo chamada dentro de "
    "um menu interativo."
)

doc.add_heading("2.5 Tratamento de erros como parte do raciocínio lógico", level=2)
doc.add_paragraph(
    "Um sistema de produção real recebe entradas inesperadas: texto onde se "
    "espera número, um id de peça repetido, uma tentativa de remover uma peça "
    "que já foi expedida numa caixa fechada. Cada uma dessas situações foi "
    "modelada como uma exceção de domínio específica (PecaDuplicadaError, "
    "PecaNaoEncontradaError, RemocaoBloqueadaError) em vez de deixar o programa "
    "quebrar com uma mensagem genérica do Python. O menu captura cada uma "
    "dessas exceções e traduz para uma mensagem clara ao operador, sem nunca "
    "derrubar o programa."
)

# ============================== 3. BENEFÍCIOS ==============================
doc.add_heading("3. Benefícios percebidos na solução", level=1)
beneficios = [
    (
        "Consistência absoluta de julgamento",
        "Uma peça de 105.0g é sempre aprovada; uma de 105.1g é sempre reprovada. "
        "Não há variação entre operadores, turnos ou dias.",
    ),
    (
        "Rastreabilidade completa",
        "Cada peça carrega o motivo exato da sua reprovação e a caixa exata onde "
        "foi armazenada, se aprovada — persistindo entre execuções do programa.",
    ),
    (
        "Feedback imediato e completo",
        "O operador sabe, no instante do cadastro, se a peça foi aprovada e por "
        "quê — sem esperar um relatório de fim de turno.",
    ),
    (
        "Auditabilidade do lote",
        "A imutabilidade das caixas fechadas garante que, uma vez lacrado um lote "
        "de 10 peças, seu conteúdo nunca muda silenciosamente.",
    ),
    (
        "Confiabilidade comprovada por testes",
        "21 testes automatizados travam o comportamento esperado em casos de "
        "borda — qualquer alteração futura que quebre uma regra é detectada "
        "instantaneamente.",
    ),
]
for titulo_b, texto_b in beneficios:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(titulo_b + ": ")
    run.bold = True
    p.add_run(texto_b)

# ============================== 4. DESAFIOS ==============================
doc.add_heading("4. Desafios enfrentados no desenvolvimento", level=1)
desafios = [
    (
        "Limites das faixas (inclusivo vs. exclusivo)",
        'O enunciado diz "peso entre 95g e 105g" sem especificar se os limites '
        "exatos são aprovados. Adotamos a leitura inclusiva, documentada "
        "explicitamente no código e neste relatório, em vez de resolvida "
        "silenciosamente.",
    ),
    (
        "Remoção de peça em caixa já fechada",
        "Decidimos que caixas fechadas são imutáveis — espelhando a prática real "
        "de lotes lacrados — e implementamos essa restrição como uma exceção de "
        "domínio explícita (RemocaoBloqueadaError).",
    ),
    (
        'Erro de "off-by-one" no limite de 10 peças',
        "Clássico risco em lógica de agrupamento por capacidade. Validado por "
        "testes automatizados dedicados exatamente à 10ª e à 11ª peça.",
    ),
    (
        "Robustez na leitura de entrada do terminal",
        "A função ler_float() trata vírgula decimal e repete a pergunta em caso "
        "de entrada inválida, sem exigir do operador conhecimento de formatação.",
    ),
]
for titulo_d, texto_d in desafios:
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(titulo_d + ": ")
    run.bold = True
    p.add_run(texto_d)

# ============================== 5. REFLEXÃO FINAL ==============================
doc.add_heading("5. Reflexão final: expandindo o protótipo para um cenário real", level=1)
doc.add_paragraph(
    "O protótipo em Python resolve o problema em escala de demonstração — um "
    "operador, um terminal, algumas dezenas de peças por sessão. Levá-lo a um "
    "cenário real de produção industrial contínua exigiria evoluir três "
    "frentes, sem necessariamente reescrever a lógica de negócio já "
    "implementada e testada."
)

doc.add_heading("Sensores substituindo a entrada manual", level=3)
doc.add_paragraph(
    "Uma balança digital com saída serial/Modbus e um sensor a laser de "
    "comprimento poderiam alimentar GerenciadorProducao.cadastrar_peca() "
    "automaticamente, a cada peça que passa por um ponto fixo da esteira. A "
    "cor poderia vir de uma câmera com um classificador simples de visão "
    "computacional (ex.: OpenCV), substituindo o input() de cor por uma "
    "leitura automática em milissegundos."
)

doc.add_heading("Inteligência artificial aplicada à causa raiz e à manutenção preditiva", level=3)
doc.add_paragraph(
    "O relatório atual já agrega motivos de reprovação; um passo natural "
    "seguinte é aplicar um modelo estatístico ou de machine learning sobre "
    "esse histórico para detectar padrões antes que se tornem problemas "
    "visíveis — por exemplo, se o percentual de peças reprovadas por "
    '"comprimento fora da faixa" começa a subir de forma consistente ao '
    "longo de um turno, isso pode indicar desgaste de uma ferramenta de corte "
    "muito antes que o problema afete um lote inteiro."
)

doc.add_heading("Integração industrial (MES/ERP e atuação física)", level=3)
doc.add_paragraph(
    'Num cenário real, a decisão "reprovada" deveria acionar um atuador '
    "físico (ex.: um desvio pneumático na esteira) que remove automaticamente "
    "a peça defeituosa da linha principal, e o sistema deveria se comunicar "
    "com um MES/ERP da fábrica para que volume de produção, refugo e "
    "rendimento apareçam em dashboards gerenciais em tempo real."
)

doc.add_paragraph(
    "O ponto que sustenta toda essa reflexão é arquitetural: como o código já "
    'separa claramente "como os dados chegam", "o que é decidido" e '
    '"onde os dados vão", cada uma dessas evoluções — sensores, IA, '
    "integração industrial — substitui apenas as bordas do sistema. A lógica "
    "de decisão que hoje aprova ou reprova uma peça, e que hoje fecha uma caixa "
    "ao atingir 10 unidades, permaneceria idêntica e já validada por 21 testes "
    "automatizados, mesmo que a peça chegasse por um sensor IoT em vez de por "
    "um teclado."
)

# ============================== REFERÊNCIAS ==============================
doc.add_heading("Referências", level=1)
for ref in [
    "Documentação oficial da linguagem Python: https://docs.python.org/3/",
    "Python Brasil — comunidade e recursos em português: https://python.org.br/",
    "Material da disciplina de Algoritmos e Lógica de Programação — UniFECAF.",
]:
    doc.add_paragraph(ref, style="List Bullet")

doc.save("RELATORIO_TECNICO.docx")
print("RELATORIO_TECNICO.docx gerado com sucesso.")
